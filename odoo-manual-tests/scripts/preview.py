#!/usr/bin/env python3
"""Render each sheet of a checklist workbook to a PNG for a visual check (macOS Quick Look).

Usage:
    python3 preview.py BOOK.xlsx --outdir DIR [--sheet NAME ...] [--from-row N]

Quick Look only draws the first sheet, and crops sheets that are much wider than
tall. The script therefore renders one temporary copy per sheet with that sheet
first and a tall blank row appended; only the PNGs are left in DIR. --from-row
hides the rows above N to look further down a long sheet. Quick Look ignores
bold runs and shows formulas as 0 — both are fine in Excel and Numbers.
"""
import argparse
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

MIN_HEIGHT_PT = 1100  # below this, Quick Look crops wide sheets on the right


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("book")
    parser.add_argument("--outdir", required=True)
    parser.add_argument("--sheet", action="append", help="sheet name (repeatable); default: all")
    parser.add_argument("--from-row", type=int, default=0, help="hide rows 2..N-1 on the chosen sheets")
    args = parser.parse_args()
    if not shutil.which("qlmanage"):
        sys.exit("qlmanage not found (macOS only). Open the file in Excel or Numbers instead.")
    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    names = args.sheet or load_workbook(args.book, read_only=True).sheetnames
    with tempfile.TemporaryDirectory(prefix="checklist_preview_") as tmp:
        for name in names:
            wb = load_workbook(args.book, rich_text=True)
            if name not in wb.sheetnames:
                print(f"skip: no sheet {name!r}")
                continue
            wb.move_sheet(name, offset=-wb.sheetnames.index(name))
            wb.active = 0
            for ws in wb.worksheets:
                ws.sheet_view.tabSelected = ws.title == name
            ws = wb[name]
            if args.from_row > 2:
                for r in range(2, args.from_row):
                    ws.row_dimensions[r].hidden = True
            height = sum((ws.row_dimensions[r].height or 15) for r in range(1, ws.max_row + 1)
                         if not ws.row_dimensions[r].hidden)
            pad_row = ws.max_row + 1
            while height < MIN_HEIGHT_PT:
                ws.row_dimensions[pad_row].height = 400
                ws.cell(pad_row, 1, " ")
                height += 400
                pad_row += 1
            copy = Path(tmp) / f"{Path(args.book).stem}__{name.replace(' ', '_')}.xlsx"
            wb.save(copy)
            subprocess.run(["qlmanage", "-t", "-s", "2400", "-o", str(outdir), str(copy)],
                           capture_output=True, check=False)
            png = outdir / (copy.name + ".png")
            print(png if png.exists() else f"no image for {name!r}")


if __name__ == "__main__":
    main()
