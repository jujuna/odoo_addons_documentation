#!/usr/bin/env python3
"""Build a manual test checklist workbook (xlsx) from a JSON spec.

Usage:
    python3 build_checklist.py SPEC.json [--out PATH]

The spec format is documented in ../references/format-rules.md; a full worked
example is ../assets/example_rs_einvoice.json. The script also lints the spec
(unknown "Look here" names, missing starting points, long cells, case count)
and prints the warnings before saving.
"""
import argparse
import json
import math
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
NAVY = "FF1F3864"
BLUE = "FF2F5597"
GREEN_DARK = "FF38761D"
ORANGE = "FFC55A11"
GREY_BAR = "FF404040"
THIN = Side(style="thin", color="FFBFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
WRAP_MID = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(wrap_text=True, vertical="top", horizontal="center")
CENTER_MID = Alignment(wrap_text=True, vertical="center", horizontal="center")
TOKEN = re.compile(r"(\[[^\]]+\]|\*[^*]+\*)")
# "Look here" values that need no entry in the where table.
FREE_LOOK = {"—", "whole screen", "the list that opens", "table on that page", "tabs of the user page"}


def font(size=10, bold=False, italic=False, color="FF000000"):
    return Font(name=FONT, size=size, bold=bold, italic=italic, color=color)


def fill(color):
    return PatternFill("solid", fgColor=color)


def rich(text, size=10, color="FF000000"):
    """Return rich text: [Button] bold with brackets, *value* bold without asterisks."""
    if not isinstance(text, str) or not TOKEN.search(text):
        return text
    base = InlineFont(rFont=FONT, sz=size, color=color)
    bold = InlineFont(rFont=FONT, sz=size, color=color, b=True)
    blocks = CellRichText()
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("[") and part.endswith("]"):
            blocks.append(TextBlock(bold, part))
        elif part.startswith("*") and part.endswith("*"):
            blocks.append(TextBlock(bold, part[1:-1]))
        else:
            blocks.append(TextBlock(base, part))
    return blocks


def n_lines(text, width):
    per_line = max(int(width * 1.1), 6)
    text = (text or "").replace("*", "")
    return sum(max(1, math.ceil(len(p) / per_line)) for p in text.split("\n"))


def fit(ws, row, pairs, minimum=17):
    """Set a row height that shows every wrapped line (Excel does not auto-fit)."""
    ws.row_dimensions[row].height = max(minimum, max(n_lines(t, w) for t, w in pairs) * 13.5 + 6)


# ----------------------------------------------------------------------------- lint
def lint(spec):
    """Return a list of warnings about the spec; the build still runs."""
    warnings = []
    where_names = {row[0].strip().lower() for row in spec.get("where", [])}
    external = (spec.get("external") or {}).get("name", "")
    free = FREE_LOOK | ({f"{external} website".lower()} if external else set())
    points = {p[0] for p in spec.get("starting_points", [])}
    cases = [c for part in spec["parts"] for c in part["cases"]]
    safe = [c for c in cases if c.get("kind") == "safe"]
    for n, case in enumerate(cases, start=1):
        tag = f"TC-{n:02d} {case['title']!r}"
        if case.get("kind") not in ("flow", "safe"):
            warnings.append(f"{tag}: kind must be 'flow' or 'safe'")
        start = (case.get("start") or "—").strip()
        first = start.split(",")[0].split(" ")[0]
        if len(first) == 1 and first.isalpha() and first not in points:
            warnings.append(f"{tag}: starting point {first!r} is not defined in starting_points")
        if not case.get("steps"):
            warnings.append(f"{tag}: no steps")
        for i, step in enumerate(case.get("steps", []), start=1):
            if len(step) != 4:
                warnings.append(f"{tag} step {i}: needs [where, do, look, see]")
                continue
            where, do, look, see = step
            if where not in ("odoo", "ext"):
                warnings.append(f"{tag} step {i}: where must be 'odoo' or 'ext'")
            if where == "ext" and not external:
                warnings.append(f"{tag} step {i}: 'ext' step but the spec has no 'external'")
            for piece in (p.strip() for p in look.split(" + ")):
                if piece.lower() not in where_names and piece.lower() not in free:
                    warnings.append(f"{tag} step {i}: Look here {piece!r} is not a name in the 'where' table")
            if len(do.replace("*", "")) > 230:
                warnings.append(f"{tag} step {i}: 'Do this' is long ({len(do)} chars) — split the step")
            if len(see.replace("*", "")) > 130:
                warnings.append(f"{tag} step {i}: 'You should see' is long ({len(see)} chars) — keep one check")
    target = spec.get("target_cases")
    if target and not (0.8 * target <= len(cases) <= 1.2 * target):
        warnings.append(f"{len(cases)} test cases, target is {target}")
    if cases and not (0.15 <= len(safe) / len(cases) <= 0.4):
        warnings.append(f"{len(safe)} of {len(cases)} cases are safety cases — aim for 20-30%")
    return warnings


# ---------------------------------------------------------------------------- sheets
def build_start(wb, spec, n_cases):
    ws = wb.active
    ws.title = "Start"
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABC", (26, 22, 80)):
        ws.column_dimensions[col].width = w
    state = {"r": 1}
    ws.cell(1, 1, spec["title"]).font = font(18, True, color=NAVY)
    ws.row_dimensions[1].height = 30
    state["r"] = 2
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=3)
    intro = spec.get("intro") or (
        f"{n_cases} test cases for this module. You don't need to know Odoo: each row says "
        "what to click, where to look and what you must see.")
    intro = intro.replace("{n_cases}", str(n_cases))
    c = ws.cell(2, 1, intro)
    c.font = font(11)
    c.alignment = WRAP
    ws.row_dimensions[2].height = max(34, n_lines(intro, 128) * 15 + 6)

    def head(title):
        state["r"] += 2
        r = state["r"]
        ws.cell(r, 1, title).font = font(12, True, color=NAVY)
        for col in (1, 2, 3):
            ws.cell(r, col).fill = fill("FFD9E2F3")
        ws.row_dimensions[r].height = 21

    def line(label, text, warn=False, inp=False):
        state["r"] += 1
        r = state["r"]
        a = ws.cell(r, 1, rich(label, 11))
        a.font = font(11, True)
        a.alignment = WRAP
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        b = ws.cell(r, 2, rich(text, 11))
        b.font = font(11, color="FF7F7F7F" if inp else "FF000000")
        b.alignment = WRAP
        if warn:
            for col in (1, 2, 3):
                ws.cell(r, col).fill = fill("FFFDE9D9")
        if inp:
            for col in (2, 3):
                ws.cell(r, col).fill = fill("FFFFF2CC")
                ws.cell(r, col).border = BORDER
        fit(ws, r, [(label, 26), (text, 102)], minimum=19)

    external = spec.get("external")
    how = spec.get("how_to_use") or [
        "Read sheet *Where to find things* once (1 minute).",
        "Sheet *Tests*: do the test cases in order. Each one starts from a *starting point* (below).",
        "After a test case, pick *Pass* or *Fail* in column *Result* on its dark title row. "
        "On Fail, write the step number and what you saw in *Note*.",
        ("*Orange* rows are done on the {ext}. ".format(ext=external["label"]) if external else "")
        + "*Green* test cases check that Odoo stops you: the error message (or a button that is not "
        "there) is the Pass.",
    ]
    head("How to use it")
    for i, text in enumerate(how, start=1):
        line(str(i), text)
    if spec.get("admin_fill"):
        head("The admin fills this in before testing")
        for label, hint in spec["admin_fill"]:
            line(label, hint, inp=True)
    if spec.get("test_data"):
        head("Test data (the admin creates it once)")
        for label, text in spec["test_data"]:
            line(label, text)
    if spec.get("starting_points"):
        head("Starting points")
        for key, name, how_text in spec["starting_points"]:
            state["r"] += 1
            r = state["r"]
            a = ws.cell(r, 1, key)
            a.font = font(12, True, color=BLUE)
            a.alignment = CENTER
            b = ws.cell(r, 2, name)
            b.font = font(11, True)
            b.alignment = WRAP
            cc = ws.cell(r, 3, rich(how_text, 11))
            cc.font = font(11)
            cc.alignment = WRAP
            fit(ws, r, [(name, 22), (how_text, 80)], minimum=19)
        for label, text in spec.get("tips", []):
            line(label, text)
    if spec.get("never"):
        head("Never")
        for label, text in spec["never"]:
            line(label, text, warn=True)
    head("Progress")
    first = state["r"] + 1
    for label, word in (("Pass", "Pass"), ("Fail", "Fail")):
        state["r"] += 1
        r = state["r"]
        ws.cell(r, 1, label).font = font(11, True)
        cc = ws.cell(r, 2, f'=COUNTIF(Tests!$F$2:$F$2000,"{word}")')
        cc.font = font(11, True)
        cc.alignment = Alignment(horizontal="left")
    state["r"] += 1
    r = state["r"]
    ws.cell(r, 1, "Not done").font = font(11, True)
    cc = ws.cell(r, 2, f"={n_cases}-B{first}-B{first + 1}")
    cc.font = font(11, True)
    cc.alignment = Alignment(horizontal="left")
    ws.cell(r, 3, f"(out of {n_cases} test cases)").font = font(10, italic=True, color="FF7F7F7F")


def build_where(wb, spec):
    ww = wb.create_sheet("Where to find things")
    ww.sheet_view.showGridLines = False
    for col, w in zip("ABC", (32, 58, 62)):
        ww.column_dimensions[col].width = w
    ww.cell(1, 1, "Where to find things").font = font(16, True, color=NAVY)
    ww.row_dimensions[1].height = 26
    ww.cell(2, 1, "Column \"Look here\" on sheet Tests uses the names in the first column.").font = font(
        10, italic=True, color="FF595959")
    for col, h in enumerate(("Name used in Tests", "Where it is on the screen", "What you see there"), start=1):
        x = ww.cell(4, col, h)
        x.font = font(10, True, color="FFFFFFFF")
        x.fill = fill(NAVY)
        x.border = BORDER
        x.alignment = WRAP_MID
    ww.row_dimensions[4].height = 20
    r = 5
    for name, where, what in spec["where"]:
        for col, v in enumerate((name, where, what), start=1):
            x = ww.cell(r, col, rich(v, 11))
            x.font = font(11, col == 1)
            x.alignment = WRAP
            x.border = BORDER
            if col == 1:
                x.fill = fill("FFF2F2F2")
        fit(ww, r, [(name, 32), (where, 58), (what, 62)], minimum=20)
        r += 1
    ww.freeze_panes = "A5"


def build_tests(wb, spec):
    wt = wb.create_sheet("Tests")
    heads = ["Test", "Where", "Do this", "Look here", "You should see", "Result", "Note"]
    widths = [8, 7, 56, 26, 42, 9, 24]
    for i, (h, w) in enumerate(zip(heads, widths), start=1):
        x = wt.cell(1, i, h)
        x.font = font(10, True, color="FFFFFFFF")
        x.fill = fill(NAVY)
        x.alignment = CENTER_MID
        x.border = BORDER
        wt.column_dimensions[x.column_letter].width = w
    wt.row_dimensions[1].height = 20
    external = spec.get("external") or {}
    r, number, result_cells = 2, 0, []
    for part in spec["parts"]:
        wt.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        x = wt.cell(r, 1, part["title"])
        x.font = font(12, True, color="FFFFFFFF")
        x.alignment = WRAP_MID
        for col in range(1, 8):
            wt.cell(r, col).fill = fill(GREY_BAR)
        wt.row_dimensions[r].height = 26
        r += 1
        role = part.get("role")
        for case in part["cases"]:
            number += 1
            band = GREEN_DARK if case["kind"] == "safe" else BLUE
            for col in range(1, 8):
                wt.cell(r, col).fill = fill(band)
                wt.cell(r, col).border = BORDER
            x = wt.cell(r, 1, f"TC-{number:02d}")
            x.font = font(10, True, color="FFFFFFFF")
            x.alignment = CENTER_MID
            wt.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
            x = wt.cell(r, 2, case["title"])
            x.font = font(11, True, color="FFFFFFFF")
            x.alignment = WRAP_MID
            wt.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
            start = (case.get("start") or "—").strip()
            info = f"Start: {start}\n{case['goal']}" if start != "—" else case["goal"]
            x = wt.cell(r, 4, info)
            x.font = font(10, italic=True, color="FFFFFFFF")
            x.alignment = WRAP_MID
            res = wt.cell(r, 6)
            res.fill = fill("FFFFF2CC")
            res.alignment = CENTER_MID
            res.font = font(10, True)
            result_cells.append(res.coordinate)
            wt.cell(r, 7).fill = fill("FFFFFFFF")
            wt.row_dimensions[r].height = max(28, max(n_lines(info, 68), n_lines(case["title"], 63)) * 13.5 + 8)
            r += 1
            for i, (where, do, look, see) in enumerate(case["steps"], start=1):
                is_ext = where == "ext"
                if is_ext and role:
                    do = f"*As {role}:* {do}"
                label = external.get("name", "Web") if is_ext else "Odoo"
                values = [str(i), label, rich(do), look, rich(see), None, None]
                for col, v in enumerate(values, start=1):
                    x = wt.cell(r, col, v)
                    x.border = BORDER
                    x.font = font(10)
                    x.alignment = CENTER if col in (1, 2) else WRAP
                wt.cell(r, 2).font = font(10, True, color=ORANGE if is_ext else BLUE)
                wt.cell(r, 4).font = font(10, color="FF404040")
                row_fill = "FFFBE5D6" if is_ext else ("FFE2EFDA" if case["kind"] == "safe" else None)
                if row_fill:
                    for col in range(1, 6):
                        wt.cell(r, col).fill = fill(row_fill)
                fit(wt, r, [(do, 56), (look, 26), (see, 42)])
                r += 1
    last = r - 1
    dv = DataValidation(type="list", formula1='"Pass,Fail"', allow_blank=True)
    wt.add_data_validation(dv)
    for coord in result_cells:
        dv.add(coord)
    for text, color in (("Pass", "FFC6EFCE"), ("Fail", "FFFFC7CE")):
        wt.conditional_formatting.add(
            f"F2:F{last}",
            CellIsRule(operator="equal", formula=[f'"{text}"'],
                       fill=PatternFill("solid", fgColor=color, bgColor=color)),
        )
    wt.freeze_panes = "C2"
    wt.page_setup.orientation = "landscape"
    wt.page_setup.fitToWidth = 1
    wt.page_setup.fitToHeight = 0
    wt.sheet_properties.pageSetUpPr.fitToPage = True
    wt.print_title_rows = "1:1"
    return number


def build_cheat(wb, spec):
    if not spec.get("cheat_sheet"):
        return
    wc = wb.create_sheet("Cheat sheet")
    wc.sheet_view.showGridLines = False
    for col, w in zip("ABC", (26, 50, 50)):
        wc.column_dimensions[col].width = w
    r = 1
    for table in spec["cheat_sheet"]:
        wc.cell(r, 1, table["title"]).font = font(12, True, color=NAVY)
        r += 1
        head = [h for h in table["head"] if h]
        for col, h in enumerate(head, start=1):
            y = wc.cell(r, col, h)
            y.font = font(10, True, color="FFFFFFFF")
            y.fill = fill(NAVY)
            y.border = BORDER
        r += 1
        for row in table["rows"]:
            row = row[:len(head)]
            for col, v in enumerate(row, start=1):
                y = wc.cell(r, col, rich(v))
                y.font = font(10, col == 1)
                y.alignment = WRAP
                y.border = BORDER
            fit(wc, r, list(zip(row, (26, 50, 50))))
            r += 1
        r += 1


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("spec")
    parser.add_argument("--out", help="output xlsx (default: spec['output'])")
    parser.add_argument("--check", action="store_true", help="lint only, write nothing (exit 1 on warnings)")
    parser.add_argument("--force", action="store_true", help="overwrite an existing file (never one open in Excel)")
    args = parser.parse_args()
    spec = json.loads(Path(args.spec).read_text(encoding="utf-8"))
    warnings = lint(spec)
    for w in warnings:
        print("WARNING:", w)
    if args.check:
        n = sum(len(part["cases"]) for part in spec["parts"])
        print(f"check: {n} test cases, {len(warnings)} warning(s), nothing written")
        sys.exit(1 if warnings else 0)
    out = args.out or spec.get("output")
    if not out:
        sys.exit("No output path: pass --out or set 'output' in the spec.")
    lock = Path(out).with_name("~$" + Path(out).name)
    if lock.exists():
        sys.exit(f"{out} is open in Excel ({lock.name} exists) and may hold the tester's results. "
                 "Write to a new file name instead.")
    if Path(out).exists() and not args.force:
        sys.exit(f"{out} already exists and may hold the tester's results. "
                 "Check it, then pass --force to overwrite, or use a new file name.")
    cases = [c for part in spec["parts"] for c in part["cases"]]
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    build_start(wb, spec, len(cases))
    build_where(wb, spec)
    build_tests(wb, spec)
    build_cheat(wb, spec)
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    safe = sum(1 for c in cases if c["kind"] == "safe")
    steps = sum(len(c["steps"]) for c in cases)
    print(f"saved {out} | {len(cases)} test cases ({len(cases) - safe} flow, {safe} safety) | "
          f"{steps} steps | {len(warnings)} warning(s)")


if __name__ == "__main__":
    main()
